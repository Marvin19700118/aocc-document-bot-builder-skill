import json
from pathlib import Path
import subprocess
import sys
import time

import pytest

from conftest import ROOT, TestTokenizer
from docbot_core.parsing import Passage, chunk, parse
from docbot_core.storage import BotError, FileLock
from install import install


def pdf_bytes(text="The warranty period is 24 months."):
    # Minimal valid one-page PDF fixture, not an image/OCR mock.
    stream = f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode("ascii")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>", b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>", b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    result, offsets = b"%PDF-1.4\n", [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(result)
    result += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets[1:])
    result += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return result


def test_pdf_real_extraction_and_page(tmp_path):
    path = tmp_path / "warranty.pdf"
    path.write_bytes(pdf_bytes())
    passages, warnings = parse(path)
    assert "24 months" in passages[0].text
    assert passages[0].locator["page"] == 1
    assert warnings == []


def test_docx_paragraph_table_and_no_fabricated_page(tmp_path):
    from docx import Document
    path = tmp_path / "制度.docx"
    document = Document()
    document.add_heading("公司制度", level=1)
    document.add_paragraph("保固期限為兩年。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "產品", "Aurora"
    document.save(path)
    passages, warnings = parse(path)
    assert any(p.locator.get("paragraph") == 2 and "兩年" in p.text for p in passages)
    assert any(p.locator.get("table") == 1 and p.locator["row"] == 1 for p in passages)
    assert not any("page" in p.locator for p in passages)


def test_markdown_text_locations_and_chunk_offsets(tmp_path):
    path = tmp_path / "手冊.md"
    path.write_text("# 說明\n\n第一行\n第二行\n", encoding="utf-8")
    passages, _ = parse(path)
    assert passages[-1].locator["line_start"] == 3
    assert passages[-1].locator["line_end"] == 4
    long = Passage("x"*950, {"type": "text", "line_start": 1, "line_end": 1})
    chunks = chunk([long], TestTokenizer())
    assert [len(c.text) for c in chunks] == [400, 400, 270]
    assert chunks[1].locator["char_start"] == 340


def test_scanned_and_encrypted_pdf_explicitly_rejected(tmp_path):
    from pypdf import PdfWriter
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    with pytest.raises(BotError) as error:
        parse(path)
    assert error.value.code == "no_extractable_text"
    writer.encrypt("test")
    writer.write(path)
    with pytest.raises(BotError) as error:
        parse(path)
    assert error.value.code == "encrypted_pdf"


def test_install_two_hosts_and_update_owned_only(tmp_path):
    for host in ("claude", "codex"):
        result = install(host, tmp_path)
        target = Path(result["path"])
        assert (target / "SKILL.md").is_file()
        assert (target / "scripts" / "docbot_core" / "engine.py").is_file()
        with pytest.raises(ValueError):
            install(host, tmp_path)
        install(host, tmp_path, update=True)
    unrelated = tmp_path / ".claude" / "skills" / "another-skill"
    unrelated.mkdir()
    (unrelated / "keep").write_text("keep")
    install("claude", tmp_path, update=True)
    assert (unrelated / "keep").read_text() == "keep"


def test_installer_refuses_unowned_same_name_even_with_update(tmp_path):
    target = tmp_path / ".agents" / "skills" / "document-bot"
    target.mkdir(parents=True)
    original = target / "SKILL.md"
    original.write_text("an unrelated skill with the same name")
    with pytest.raises(ValueError):
        install("codex", tmp_path, update=True)
    assert original.read_text() == "an unrelated skill with the same name"


def test_cli_json_status_and_empty_evidence(bot):
    script = ROOT / "document-bot" / "scripts" / "document_bot.py"
    result = subprocess.run([sys.executable, str(script), "--home", str(bot.store.home), "status"], capture_output=True, text=True, encoding="utf-8", timeout=30)
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["data"]["graph_enabled"] is True


def test_lock_excludes_separate_process_and_crash_releases(tmp_path):
    lock = tmp_path / "worker.lock"
    ready = tmp_path / "ready"
    script = "from pathlib import Path; import sys,time; sys.path.insert(0,sys.argv[1]); from docbot_core.storage import FileLock; lock=FileLock(sys.argv[2]); lock.__enter__(); Path(sys.argv[3]).write_text('ready'); time.sleep(30)"
    child = subprocess.Popen([sys.executable, "-c", script, str(ROOT / "document-bot" / "scripts"), str(lock), str(ready)])
    try:
        deadline = time.time()+10
        while not ready.exists() and time.time() < deadline:
            time.sleep(.05)
        assert ready.exists()
        with pytest.raises(BotError):
            with FileLock(lock):
                pass
    finally:
        child.terminate()
        child.wait(timeout=10)
    with FileLock(lock):
        pass
