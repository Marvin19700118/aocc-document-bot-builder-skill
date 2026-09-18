"""Real-model retrieval check. Generated fixtures are not claimed as host uploads."""
import argparse
import json
from pathlib import Path
import platform
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "document-bot" / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

from docbot_core.engine import Bot
from test_parsing_install import pdf_bytes


QUESTIONS = [
    ("特休假每年有幾天？", "20 天", "員工手冊.md"),
    ("遠端工作每週最多幾天？", "3 天", "員工手冊.md"),
    ("出差住宿每晚補助上限？", "3,000 元", "員工手冊.md"),
    ("教育訓練每年的預算？", "18,000 元", "員工手冊.md"),
    ("離職前需提前幾天通知？", "30 天", "員工手冊.md"),
    ("Aurora 的資料匯出格式？", "JSON 與 CSV", "產品規格.docx"),
    ("Aurora 客服電話是什麼？", "02-5555-0100", "產品規格.docx"),
    ("How long is the hardware warranty?", "24 months", "warranty.pdf"),
    ("備份資料保存多久？", "90 天", "維運說明.txt"),
    ("維護時段是哪一天幾點？", "星期日凌晨 02:00", "維運說明.txt"),
    ("董事長的私人手機號碼？", None, None),
    ("公司明年的營收預測？", None, None),
    ("Aurora 的上市股價是多少？", None, None),
    ("公司醫療保險的理賠比例？", None, None),
    ("產品是否通過 ISO 27001 認證？", None, None),
]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-download", action="store_true")
    args = parser.parse_args()
    from docx import Document
    import psutil
    base = ROOT / ".test-runtime" / "live"
    workspace = base / "workspace"
    source = base / "generated-fixtures"
    workspace.mkdir(parents=True, exist_ok=True)
    source.mkdir(parents=True, exist_ok=True)
    (source / "員工手冊.md").write_text("# 員工手冊\n\n特休假每年 20 天。\n\n遠端工作每週最多 3 天。\n\n出差住宿每晚補助上限新台幣 3,000 元。\n\n教育訓練每年預算新台幣 18,000 元。\n\n離職前需提前 30 天通知。\n\n" + "\n\n".join(f"行政流程 {i}：會議室設備使用完畢，應關閉電源並歸還借用物品。" for i in range(1, 21)), encoding="utf-8")
    document = Document()
    document.add_heading("Aurora 產品規格", level=1)
    document.add_paragraph("星河公司開發 Aurora。Aurora 的資料匯出格式為 JSON 與 CSV。")
    document.add_paragraph("Aurora 客服電話為 02-5555-0100。")
    document.save(source / "產品規格.docx")
    (source / "warranty.pdf").write_bytes(pdf_bytes())
    (source / "維運說明.txt").write_text("備份資料保存 90 天。\n\n維護時段是星期日凌晨 02:00。\n", encoding="utf-8")
    bot = Bot(base / "home", initialize=True)
    started = time.perf_counter()
    if args.skip_download:
        bot.setup(workspace, download=False)
    else:
        bot.setup(workspace, download=True)
    setup_seconds = time.perf_counter()-started
    # Vector-only baseline. Graph semantic checks are performed separately by the assistant.
    bot.graph_toggle(False)
    manifest = {"source": "chat_attachment", "host": "codex", "order_known": True, "attachments": [{"attachment_id": "generated-fixture:"+p.name, "name": p.name, "path": str(p.resolve())} for p in sorted(source.iterdir())]}
    receipts = bot.add(manifest)
    started = time.perf_counter()
    work = bot.work()
    indexing_seconds = time.perf_counter()-started
    results = []
    for question, expected, filename in QUESTIONS:
        started = time.perf_counter()
        retrieved = bot.ask(question)
        results.append({"question": question, "expected_excerpt": expected, "expected_file": filename, "retrieval_seconds": time.perf_counter()-started, "expected_evidence_found": any(expected in e["text"] and e["filename"] == filename for e in retrieved["evidence"]) if expected else None, "retrieved": retrieved})
    report = {"platform": platform.platform(), "python": platform.python_version(), "source_adapter": "generated fixture manifest; not a real chat attachment test", "model": bot.store.setting("model"), "model_revision": bot.store.setting("model_revision"), "setup_seconds": setup_seconds, "indexing_seconds": indexing_seconds, "rss_bytes": psutil.Process().memory_info().rss, "receipts": receipts, "work": work, "documents": bot.documents(), "questions": results}
    output = ROOT / "verification" / "live-retrieval.json"
    output.parent.mkdir(exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(output), "answerable_retrieval_hits": sum(r["expected_evidence_found"] is True for r in results), "answerable_total": 10, "unanswerable": "requires host answer review; retrieval alone does not prove abstention", "indexing_seconds": indexing_seconds, "rss_bytes": report["rss_bytes"]}, ensure_ascii=False))
    bot.store.close()


if __name__ == "__main__":
    main()
